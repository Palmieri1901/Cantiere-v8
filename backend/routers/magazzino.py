"""Endpoints magazzino: articoli, fornitori, movimenti, scan AI, export PDF/Excel."""
import io
import os
import base64
import json
import re
from datetime import datetime, timezone
from typing import List, Optional
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse

from database import db
from models import (
    Articolo, ArticoloCreate,
    Fornitore, FornitoreCreate,
    MovimentoMagazzino, MovimentoCreate,
    ScanArticoloRequest, ScanDDTRequest,
    RicaricoCategoria, RicaricoCategoriaCreate,
    SpesaAccessoria, SpesaCreate,
)


def _to_dt(d: dict, keys=("created_at", "updated_at")):
    """Converte in datetime le chiavi ISO string (per rifare passare Articolo/Fornitore)."""
    for k in keys:
        v = d.get(k)
        if isinstance(v, str):
            try:
                d[k] = datetime.fromisoformat(v.replace("Z", "+00:00"))
            except Exception:
                pass
    return d

router = APIRouter(prefix="/magazzino")


# ---------------------------------------------------------------------------
# FORNITORI
# ---------------------------------------------------------------------------

@router.get("/fornitori", response_model=List[Fornitore])
async def list_fornitori():
    docs = await db.fornitori.find({}, {"_id": 0}).sort("nome", 1).to_list(1000)
    return [Fornitore(**_to_dt(d)) for d in docs]


@router.post("/fornitori", response_model=Fornitore)
async def create_fornitore(payload: FornitoreCreate):
    if not payload.nome.strip():
        raise HTTPException(400, "Nome fornitore obbligatorio")
    f = Fornitore(**payload.model_dump())
    await db.fornitori.insert_one(f.model_dump())
    return f


@router.put("/fornitori/{fid}", response_model=Fornitore)
async def update_fornitore(fid: str, payload: FornitoreCreate):
    existing = await db.fornitori.find_one({"id": fid}, {"_id": 0})
    if not existing:
        raise HTTPException(404, "Fornitore non trovato")
    data = {**existing, **payload.model_dump(exclude_unset=True)}
    f = Fornitore(**_to_dt(data))
    await db.fornitori.update_one({"id": fid}, {"$set": f.model_dump()})
    return f


@router.delete("/fornitori/{fid}")
async def delete_fornitore(fid: str):
    r = await db.fornitori.delete_one({"id": fid})
    if r.deleted_count == 0:
        raise HTTPException(404, "Fornitore non trovato")
    # Distacca gli articoli collegati
    await db.articoli.update_many({"fornitore_id": fid}, {"$set": {"fornitore_id": None}})
    return {"ok": True}


@router.post("/fornitori/{fid}/applica-ricarico")
async def applica_ricarico_fornitore(fid: str):
    """Applica il ricarico % del fornitore a TUTTI i suoi articoli, ricalcolando
    prezzo_listino = prezzo_acquisto × (1 + ricarico/100).
    Salta gli articoli senza prezzo_acquisto valido."""
    f = await db.fornitori.find_one({"id": fid}, {"_id": 0})
    if not f:
        raise HTTPException(404, "Fornitore non trovato")
    ricarico = f.get("ricarico_default_percent")
    if ricarico is None:
        raise HTTPException(400, "Nessun ricarico % impostato per questo fornitore")
    ricarico = float(ricarico)

    docs = await db.articoli.find({"fornitore_id": fid}, {"_id": 0, "id": 1, "prezzo_acquisto": 1}).to_list(10000)
    updated = 0
    skipped = 0
    now = datetime.now(timezone.utc)
    for d in docs:
        pa = float(d.get("prezzo_acquisto") or 0)
        if pa <= 0:
            skipped += 1
            continue
        nuovo_listino = round(pa * (1 + ricarico / 100), 2)
        await db.articoli.update_one(
            {"id": d["id"]},
            {"$set": {"prezzo_listino": nuovo_listino, "updated_at": now}},
        )
        updated += 1
    return {
        "ok": True,
        "fornitore_id": fid,
        "ricarico_percent": ricarico,
        "articoli_aggiornati": updated,
        "articoli_saltati": skipped,
        "totale": len(docs),
    }


# ---------------------------------------------------------------------------
# ARTICOLI
# ---------------------------------------------------------------------------

@router.get("/articoli", response_model=List[Articolo])
async def list_articoli(
    q: Optional[str] = None,
    fornitore_id: Optional[str] = None,
    categoria: Optional[str] = None,
    sotto_scorta: Optional[bool] = False,
):
    query = {}
    if fornitore_id:
        query["fornitore_id"] = fornitore_id
    if categoria:
        query["categoria"] = categoria
    if q:
        rx = {"$regex": re.escape(q), "$options": "i"}
        query["$or"] = [{"codice": rx}, {"nome": rx}, {"descrizione": rx}]
    docs = await db.articoli.find(query, {"_id": 0}).sort("nome", 1).to_list(5000)
    articoli = [Articolo(**_to_dt(d)) for d in docs]
    if sotto_scorta:
        articoli = [a for a in articoli if a.quantita <= a.scorta_minima]
    return articoli


@router.get("/articoli/categorie", response_model=List[str])
async def list_categorie():
    cats = await db.articoli.distinct("categoria")
    return sorted([c for c in cats if c])


@router.get("/articoli/count-sotto-scorta")
async def count_sotto_scorta():
    docs = await db.articoli.find({}, {"_id": 0, "quantita": 1, "scorta_minima": 1}).to_list(5000)
    n = sum(1 for d in docs if d.get("quantita", 0) <= d.get("scorta_minima", 0))
    return {"count": n, "total": len(docs)}


@router.post("/articoli", response_model=Articolo)
async def create_articolo(payload: ArticoloCreate):
    if not payload.nome.strip():
        raise HTTPException(400, "Il nome dell'articolo è obbligatorio")
    art = Articolo(**payload.model_dump())
    await db.articoli.insert_one(art.model_dump())
    # Movimento iniziale se quantità > 0
    if art.quantita > 0:
        mv = MovimentoMagazzino(
            articolo_id=art.id, tipo="carico", quantita=art.quantita,
            quantita_dopo=art.quantita, motivo="Giacenza iniziale",
            data=datetime.now().strftime("%Y-%m-%d"),
        )
        await db.movimenti_magazzino.insert_one(mv.model_dump())
    return art


@router.put("/articoli/{aid}", response_model=Articolo)
async def update_articolo(aid: str, payload: ArticoloCreate):
    existing = await db.articoli.find_one({"id": aid}, {"_id": 0})
    if not existing:
        raise HTTPException(404, "Articolo non trovato")
    data = {**existing, **payload.model_dump(exclude_unset=True)}
    data["updated_at"] = datetime.now(timezone.utc)
    art = Articolo(**_to_dt(data))
    await db.articoli.update_one({"id": aid}, {"$set": art.model_dump()})
    return art


@router.delete("/articoli/{aid}")
async def delete_articolo(aid: str):
    r = await db.articoli.delete_one({"id": aid})
    if r.deleted_count == 0:
        raise HTTPException(404, "Articolo non trovato")
    await db.movimenti_magazzino.delete_many({"articolo_id": aid})
    return {"ok": True}


# ---------------------------------------------------------------------------
# MOVIMENTI (carico/scarico/rettifica)
# ---------------------------------------------------------------------------

@router.get("/movimenti", response_model=List[MovimentoMagazzino])
async def list_movimenti(articolo_id: Optional[str] = None, limit: int = 200):
    query = {"articolo_id": articolo_id} if articolo_id else {}
    docs = await db.movimenti_magazzino.find(query, {"_id": 0}).sort("created_at", -1).to_list(limit)
    return [MovimentoMagazzino(**_to_dt(d)) for d in docs]


@router.post("/movimenti", response_model=MovimentoMagazzino)
async def create_movimento(payload: MovimentoCreate):
    art = await db.articoli.find_one({"id": payload.articolo_id}, {"_id": 0})
    if not art:
        raise HTTPException(404, "Articolo non trovato")
    if payload.tipo not in ("carico", "scarico", "rettifica"):
        raise HTTPException(400, "Tipo movimento non valido")

    qt_prima = float(art.get("quantita", 0))
    if payload.tipo == "carico":
        nuova = qt_prima + abs(payload.quantita)
    elif payload.tipo == "scarico":
        nuova = qt_prima - abs(payload.quantita)
    else:  # rettifica: quantita è il valore assoluto finale
        nuova = float(payload.quantita)

    mv = MovimentoMagazzino(
        articolo_id=payload.articolo_id,
        tipo=payload.tipo,
        quantita=payload.quantita,
        quantita_dopo=nuova,
        motivo=payload.motivo or "",
        data=payload.data or datetime.now().strftime("%Y-%m-%d"),
        note=payload.note or "",
    )
    await db.movimenti_magazzino.insert_one(mv.model_dump())
    await db.articoli.update_one(
        {"id": payload.articolo_id},
        {"$set": {"quantita": nuova, "updated_at": datetime.now(timezone.utc)}},
    )
    return mv


# ---------------------------------------------------------------------------
# AI SCAN (Gemini 3 Flash via Emergent LLM key)
# ---------------------------------------------------------------------------

async def _run_vision(prompt: str, image_base64: str) -> str:
    from emergentintegrations.llm.chat import LlmChat, UserMessage, ImageContent, TextDelta, StreamDone

    api_key = os.environ.get("EMERGENT_LLM_KEY")
    if not api_key:
        raise HTTPException(500, "EMERGENT_LLM_KEY non configurata")

    # Rimuove eventuale prefisso data URL
    b64 = image_base64.split(",", 1)[-1] if "," in image_base64 else image_base64

    chat = (
        LlmChat(
            api_key=api_key,
            session_id=f"magazzino-{datetime.now().timestamp()}",
            system_message=(
                "Sei un assistente esperto di magazzino nautico. "
                "Analizzi foto di articoli o Documenti Di Trasporto (DDT) italiani "
                "ed estrai i dati in JSON preciso senza aggiungere commenti."
            ),
        )
        .with_model("gemini", "gemini-3-flash-preview")
    )

    image = ImageContent(image_base64=b64)
    msg = UserMessage(text=prompt, file_contents=[image])

    chunks: List[str] = []
    async for ev in chat.stream_message(msg):
        if isinstance(ev, TextDelta):
            chunks.append(ev.content)
        elif isinstance(ev, StreamDone):
            break
    return "".join(chunks)


def _extract_json(text: str):
    """Estrae il primo blocco JSON valido (oggetto o array) dal testo dell'LLM."""
    if not text:
        return None
    # Rimuovi fenced code block
    m = re.search(r"```(?:json)?\s*(.+?)```", text, re.S)
    candidate = m.group(1) if m else text
    # Trova primo { o [
    for open_ch, close_ch in (("{", "}"), ("[", "]")):
        i = candidate.find(open_ch)
        if i == -1:
            continue
        depth = 0
        for j in range(i, len(candidate)):
            if candidate[j] == open_ch:
                depth += 1
            elif candidate[j] == close_ch:
                depth -= 1
                if depth == 0:
                    try:
                        return json.loads(candidate[i:j+1])
                    except Exception:
                        break
    try:
        return json.loads(candidate.strip())
    except Exception:
        return None


@router.post("/scan-articolo")
async def scan_articolo(payload: ScanArticoloRequest):
    """Analizza la foto di un articolo/etichetta e restituisce codice, nome,
    descrizione, prezzo suggerito."""
    prompt = (
        "Analizza questa foto di un articolo nautico o della sua confezione/etichetta. "
        "Estrai i dati che vedi e rispondi SOLO con un oggetto JSON con le chiavi:\n"
        '{"codice": "", "nome": "", "descrizione": "", "categoria": "", "prezzo_listino": 0}\n\n'
        "Regole:\n"
        "- Se non riesci a leggere un campo lascialo stringa vuota (o 0 per il prezzo).\n"
        "- codice: sigla/codice prodotto stampato (es. 'RX-1234').\n"
        "- nome: nome dell'articolo (es. 'Bussola magnetica 100mm').\n"
        "- descrizione: caratteristiche principali visibili sulla confezione.\n"
        "- categoria: 1-3 parole (es. 'Ferramenta', 'Ricambi motore', 'Vernici', 'Elettronica').\n"
        "- prezzo_listino: se leggi un prezzo stampato, altrimenti 0.\n"
        "Rispondi SOLO con il JSON, nessun testo aggiuntivo."
    )
    try:
        raw = await _run_vision(prompt, payload.image_base64)
    except Exception as e:
        raise HTTPException(502, f"Errore AI vision: {e}")

    data = _extract_json(raw) or {}
    return {
        "codice": str(data.get("codice", "") or ""),
        "nome": str(data.get("nome", "") or ""),
        "descrizione": str(data.get("descrizione", "") or ""),
        "categoria": str(data.get("categoria", "") or ""),
        "prezzo_listino": float(data.get("prezzo_listino", 0) or 0),
        "raw": raw,
    }


async def _pdf_first_page_to_png_b64(pdf_b64: str) -> str:
    """Converte la prima pagina di un PDF (base64) in PNG base64 usando PyMuPDF."""
    import fitz  # pymupdf
    raw = base64.b64decode(pdf_b64.split(",", 1)[-1] if "," in pdf_b64 else pdf_b64)
    doc = fitz.open(stream=raw, filetype="pdf")
    if doc.page_count == 0:
        raise HTTPException(400, "PDF vuoto")
    page = doc.load_page(0)
    # Risoluzione medio-alta per OCR (200 DPI)
    pix = page.get_pixmap(dpi=200, alpha=False)
    png_bytes = pix.tobytes("png")
    doc.close()
    return base64.b64encode(png_bytes).decode("ascii")


@router.post("/scan-ddt")
async def scan_ddt(payload: ScanDDTRequest):
    """Analizza foto o PDF di un DDT italiano ed estrae la lista articoli.

    Accetta:
      - image_base64: immagine (JPG/PNG) del DDT
      - file_base64 + mime_type: file generico (immagine o PDF). Se PDF viene
        convertita la prima pagina in PNG prima dell'invio all'AI.
    """
    image_b64 = payload.image_base64
    if not image_b64 and payload.file_base64:
        mt = (payload.mime_type or "").lower()
        raw_b64 = payload.file_base64.split(",", 1)[-1] if "," in payload.file_base64 else payload.file_base64
        # Detection: se comincia con %PDF- oppure mime dichiarato PDF → conversione
        try:
            head = base64.b64decode(raw_b64[:8]) if raw_b64 else b""
        except Exception:
            head = b""
        if "pdf" in mt or head.startswith(b"%PDF"):
            try:
                image_b64 = await _pdf_first_page_to_png_b64(raw_b64)
            except Exception as e:
                raise HTTPException(400, f"Impossibile leggere il PDF: {e}")
        else:
            image_b64 = raw_b64
    if not image_b64:
        raise HTTPException(400, "Nessun file/immagine fornito")

    prompt = (
        "Analizza questa foto o scansione di un Documento Di Trasporto (DDT) o fattura italiana di "
        "forniture nautiche. Estrai TUTTE le righe.\n\n"
        "Distinguerai DUE gruppi di righe:\n"
        "  1) ARTICOLI (beni fisici, ricambi, materiali) → vanno in 'articoli'\n"
        "  2) SPESE ACCESSORIE (costi di trasporto, bancarie, spedizione, imballo, assicurazione, "
        "     carburante, sovrapprezzi non merceologici) → vanno in 'spese_accessorie'\n\n"
        "SIGNIFICATO COLONNE per gli articoli:\n"
        "- 'U.M.' / 'UM' / 'Unità' = unità di misura (es. pz, lt, kg, mt, mq, rotolo, cf, cad, set). "
        "  Normalizza SEMPRE in minuscolo e usa una di queste sigle: pz, lt, kg, mt, mq, rotolo, cf, cad, set, paio. "
        "  Se non riesci a leggerla usa 'pz'.\n"
        "- 'Prezzo unitario' / 'Prezzo listino' / 'PU' = prezzo unitario IVA COMPRESA (22%) riferito alla U.M.\n"
        "- 'Sconto %' / 'Sc%' = sconto in percentuale\n"
        "- 'Importo' / 'Totale' / 'Netto' = importo unitario NETTO scontato IVA ESCLUSA (prezzo di acquisto)\n"
        "- Se 'Importo' è il totale riga, dividilo per la quantità.\n\n"
        "Rispondi SOLO con un oggetto JSON:\n"
        '{"fornitore": "", "numero_ddt": "", "data": "", "iva_percent": 22,\n'
        ' "articoli": [\n'
        '   {"codice": "", "nome": "", "descrizione": "", "unita_misura": "pz", "quantita": 0,\n'
        '    "prezzo_listino_ivato": 0, "sconto_percent": 0, "importo_netto": 0}\n'
        " ],\n"
        ' "spese_accessorie": [\n'
        '   {"tipo": "trasporto", "descrizione": "Trasporto merce", "importo": 0}\n'
        " ]}\n\n"
        "TIPI DI SPESA riconosciuti (usa esattamente uno di questi valori nel campo 'tipo'):\n"
        "- 'bancarie' (spese bancarie, commissioni bonifico, incasso RIBA)\n"
        "- 'trasporto' (trasporto, corriere, consegna)\n"
        "- 'spedizione' (porto, franco destino, spedizione)\n"
        "- 'imballo' (imballaggio, cassa, pallet)\n"
        "- 'assicurazione'\n"
        "- 'carburante' (sovrapprezzo carburante, adeguamento gasolio)\n"
        "- 'altro' (qualsiasi altra spesa non merceologica)\n\n"
        "Regole:\n"
        "- l'importo delle spese è quello NETTO IVA esclusa (se leggi solo l'ivato, dividi per 1.22).\n"
        "- Non trattare sconti totali o abbuoni come spese.\n"
        "- Se non ci sono spese accessorie usa una lista vuota.\n"
        "- Ignora totali generali del documento.\n"
        "Rispondi SOLO con il JSON, nessun testo aggiuntivo."
    )
    try:
        raw = await _run_vision(prompt, image_b64)
    except Exception as e:
        raise HTTPException(502, f"Errore AI vision: {e}")

    data = _extract_json(raw) or {}
    articoli_raw = data.get("articoli") or []
    articoli = []
    iva_percent = float(data.get("iva_percent") or 22)
    for a in articoli_raw if isinstance(articoli_raw, list) else []:
        if not isinstance(a, dict):
            continue
        listino_iva = float(a.get("prezzo_listino_ivato", 0) or 0)
        sconto = float(a.get("sconto_percent", 0) or 0)
        netto = float(a.get("importo_netto", 0) or 0)

        # Se manca qualche valore, ricava dagli altri (assumendo IVA 22%)
        iva_mul = 1 + iva_percent / 100
        if netto == 0 and listino_iva > 0:
            listino_no_iva = listino_iva / iva_mul
            netto = round(listino_no_iva * (1 - sconto / 100), 4)
        elif listino_iva == 0 and netto > 0 and sconto == 0:
            listino_iva = round(netto * iva_mul, 4)
        elif listino_iva > 0 and netto > 0 and sconto == 0:
            listino_no_iva = listino_iva / iva_mul
            if listino_no_iva > 0:
                calc = (1 - netto / listino_no_iva) * 100
                if 0 <= calc <= 99:
                    sconto = round(calc, 2)

        # Retrocompatibilità: `prezzo_unitario` continua ad essere quello di acquisto (netto)
        um = str(a.get("unita_misura", "") or "pz").strip().lower() or "pz"
        # Normalizza sigle equivalenti
        um_map = {"pezzi": "pz", "n": "pz", "num": "pz", "nr": "pz", "litri": "lt", "l": "lt",
                  "chilogrammi": "kg", "chili": "kg", "metri": "mt", "m": "mt",
                  "metri quadri": "mq", "mq.": "mq", "confezione": "cf", "conf": "cf",
                  "cadauno": "cad", "cad.": "cad"}
        um = um_map.get(um, um)
        articoli.append({
            "codice": str(a.get("codice", "") or ""),
            "nome": str(a.get("nome", "") or ""),
            "descrizione": str(a.get("descrizione", "") or ""),
            "unita_misura": um,
            "quantita": float(a.get("quantita", 0) or 0),
            "prezzo_listino_ivato": listino_iva,
            "sconto_percent": sconto,
            "importo_netto": netto,
            "prezzo_unitario": netto,  # alias per compat con importa-articoli
        })
    return {
        "fornitore": str(data.get("fornitore", "") or ""),
        "numero_ddt": str(data.get("numero_ddt", "") or ""),
        "data": str(data.get("data", "") or ""),
        "iva_percent": iva_percent,
        "articoli": articoli,
        "spese_accessorie": [
            {
                "tipo": str(s.get("tipo", "altro") or "altro").lower().strip(),
                "descrizione": str(s.get("descrizione", "") or ""),
                "importo": float(s.get("importo", 0) or 0),
            }
            for s in (data.get("spese_accessorie") or [])
            if isinstance(s, dict) and (s.get("importo") or 0) > 0
        ],
        "raw": raw,
    }


class ImportDDTRequest(ScanDDTRequest.__base__ if False else object):
    pass


from pydantic import BaseModel as _BM


class ImportArticoliRequest(_BM):
    fornitore_id: Optional[str] = None
    articoli: List[dict] = []
    spese_accessorie: Optional[List[dict]] = None
    documento_ref: Optional[str] = ""  # numero DDT/fattura
    data_documento: Optional[str] = None  # ISO YYYY-MM-DD
    aggiorna_prezzi: Optional[bool] = True
    mantieni_ricarico: Optional[bool] = True


async def _default_markup_for(categoria: Optional[str]) -> Optional[float]:
    if not categoria:
        return None
    doc = await db.ricarichi_categoria.find_one({"categoria": categoria.strip()}, {"_id": 0, "ricarico_percent": 1})
    if doc and isinstance(doc.get("ricarico_percent"), (int, float)):
        return float(doc["ricarico_percent"])
    return None


async def _default_markup_for_articolo(fornitore_id: Optional[str], categoria: Optional[str]) -> Optional[float]:
    """Preferenza: ricarico del fornitore > ricarico default della categoria."""
    if fornitore_id:
        f = await db.fornitori.find_one({"id": fornitore_id}, {"_id": 0, "ricarico_default_percent": 1})
        if f and isinstance(f.get("ricarico_default_percent"), (int, float)):
            return float(f["ricarico_default_percent"])
    return await _default_markup_for(categoria)


@router.get("/ricarichi-categoria", response_model=List[RicaricoCategoria])
async def list_ricarichi():
    docs = await db.ricarichi_categoria.find({}, {"_id": 0}).sort("categoria", 1).to_list(500)
    return [RicaricoCategoria(**_to_dt(d)) for d in docs]


@router.post("/ricarichi-categoria", response_model=RicaricoCategoria)
async def upsert_ricarico(payload: RicaricoCategoriaCreate):
    cat = (payload.categoria or "").strip()
    if not cat:
        raise HTTPException(400, "Categoria obbligatoria")
    # Upsert su categoria
    existing = await db.ricarichi_categoria.find_one({"categoria": cat}, {"_id": 0})
    if existing:
        await db.ricarichi_categoria.update_one(
            {"categoria": cat}, {"$set": {"ricarico_percent": float(payload.ricarico_percent)}}
        )
        existing["ricarico_percent"] = float(payload.ricarico_percent)
        return RicaricoCategoria(**_to_dt(existing))
    r = RicaricoCategoria(categoria=cat, ricarico_percent=float(payload.ricarico_percent))
    await db.ricarichi_categoria.insert_one(r.model_dump())
    return r


@router.delete("/ricarichi-categoria/{rid}")
async def delete_ricarico(rid: str):
    res = await db.ricarichi_categoria.delete_one({"id": rid})
    if res.deleted_count == 0:
        raise HTTPException(404, "Ricarico non trovato")
    return {"ok": True}


@router.post("/importa-articoli")
async def importa_articoli(payload: ImportArticoliRequest):
    """Importa in blocco articoli dopo scan DDT.
    - Articolo esistente (match per codice): se aggiorna_prezzi=True aggiorna prezzo_acquisto
      e, se mantieni_ricarico=True, ricalcola prezzo_listino mantenendo il ricarico corrente.
    - Articolo nuovo: crea articolo; se esiste un ricarico default per la sua categoria
      calcola automaticamente prezzo_listino = prezzo_acquisto × (1 + ricarico/100).
    """
    created, updated, prezzi_aggiornati = 0, 0, 0
    spese_salvate = 0

    fornitore_nome = ""
    if payload.fornitore_id:
        f = await db.fornitori.find_one({"id": payload.fornitore_id}, {"_id": 0, "nome": 1})
        if f:
            fornitore_nome = f.get("nome", "")

    data_doc = payload.data_documento or datetime.now().strftime("%Y-%m-%d")
    for a in payload.articoli:
        codice = (a.get("codice") or "").strip()
        nome = (a.get("nome") or "").strip()
        if not nome and not codice:
            continue
        qt = float(a.get("quantita") or 0)
        prezzo = float(a.get("prezzo_unitario") or 0)

        existing = None
        if codice:
            existing = await db.articoli.find_one({"codice": codice}, {"_id": 0})
        if existing:
            update_set = {
                "quantita": float(existing.get("quantita", 0)) + qt,
                "fornitore_id": payload.fornitore_id or existing.get("fornitore_id"),
                "updated_at": datetime.now(timezone.utc),
            }
            # Se il DDT porta una U.M. valida e l'articolo non ne ha ancora una impostata, aggiornala
            um_ddt = (a.get("unita_misura") or "").strip().lower()
            if um_ddt and not (existing.get("unita_misura") or "").strip():
                update_set["unita_misura"] = um_ddt
            if prezzo > 0 and payload.aggiorna_prezzi:
                old_pa = float(existing.get("prezzo_acquisto") or 0)
                old_pv = float(existing.get("prezzo_listino") or 0)
                update_set["prezzo_acquisto"] = prezzo
                if payload.mantieni_ricarico and old_pa > 0 and old_pv > 0:
                    ricarico = (old_pv - old_pa) / old_pa
                    update_set["prezzo_listino"] = round(prezzo * (1 + ricarico), 2)
                elif payload.mantieni_ricarico:
                    # ricarico corrente non calcolabile: prova con default fornitore/categoria
                    default_mkup = await _default_markup_for_articolo(
                        existing.get("fornitore_id") or payload.fornitore_id,
                        existing.get("categoria"),
                    )
                    if default_mkup is not None:
                        update_set["prezzo_listino"] = round(prezzo * (1 + default_mkup / 100), 2)
                prezzi_aggiornati += 1
            await db.articoli.update_one({"id": existing["id"]}, {"$set": update_set})
            mv = MovimentoMagazzino(
                articolo_id=existing["id"], tipo="carico", quantita=qt,
                quantita_dopo=update_set["quantita"], motivo="Carico da DDT",
                data=datetime.now().strftime("%Y-%m-%d"),
                note=a.get("descrizione", ""),
            )
            await db.movimenti_magazzino.insert_one(mv.model_dump())
            updated += 1
        else:
            categoria = (a.get("categoria") or "").strip()
            prezzo_listino = 0.0
            if prezzo > 0:
                default_mkup = await _default_markup_for_articolo(payload.fornitore_id, categoria)
                if default_mkup is not None:
                    prezzo_listino = round(prezzo * (1 + default_mkup / 100), 2)
                else:
                    prezzo_listino = prezzo  # senza default, prezzo di vendita = acquisto
            art = Articolo(
                codice=codice, nome=nome or codice,
                descrizione=a.get("descrizione", ""),
                categoria=categoria,
                unita_misura=(a.get("unita_misura") or "pz").strip().lower() or "pz",
                fornitore_id=payload.fornitore_id,
                prezzo_acquisto=prezzo, prezzo_listino=prezzo_listino,
                quantita=qt,
            )
            await db.articoli.insert_one(art.model_dump())
            if qt > 0:
                mv = MovimentoMagazzino(
                    articolo_id=art.id, tipo="carico", quantita=qt,
                    quantita_dopo=qt, motivo="Carico da DDT",
                    data=datetime.now().strftime("%Y-%m-%d"),
                )
                await db.movimenti_magazzino.insert_one(mv.model_dump())
            created += 1
    # Salva le spese accessorie separate dal magazzino
    for sp in payload.spese_accessorie or []:
        if not isinstance(sp, dict):
            continue
        importo = float(sp.get("importo") or 0)
        if importo <= 0:
            continue
        spesa = SpesaAccessoria(
            tipo=str(sp.get("tipo") or "altro").lower().strip() or "altro",
            descrizione=str(sp.get("descrizione") or ""),
            importo=importo,
            data=data_doc,
            fornitore_id=payload.fornitore_id,
            fornitore_nome=fornitore_nome,
            documento_ref=payload.documento_ref or "",
        )
        await db.spese_accessorie.insert_one(spesa.model_dump())
        spese_salvate += 1

    return {"created": created, "updated": updated, "prezzi_aggiornati": prezzi_aggiornati, "spese_salvate": spese_salvate}


@router.get("/spese", response_model=List[SpesaAccessoria])
async def list_spese(
    tipo: Optional[str] = None,
    fornitore_id: Optional[str] = None,
    anno: Optional[int] = None,
    limit: int = 1000,
):
    query = {}
    if tipo:
        query["tipo"] = tipo
    if fornitore_id:
        query["fornitore_id"] = fornitore_id
    if anno:
        query["data"] = {"$regex": f"^{anno}-"}
    docs = await db.spese_accessorie.find(query, {"_id": 0}).sort("data", -1).to_list(limit)
    return [SpesaAccessoria(**_to_dt(d)) for d in docs]


@router.post("/spese", response_model=SpesaAccessoria)
async def create_spesa(payload: SpesaCreate):
    if payload.importo <= 0:
        raise HTTPException(400, "Importo obbligatorio")
    nome = payload.fornitore_nome or ""
    if payload.fornitore_id and not nome:
        f = await db.fornitori.find_one({"id": payload.fornitore_id}, {"_id": 0, "nome": 1})
        if f:
            nome = f.get("nome", "")
    sp = SpesaAccessoria(
        tipo=payload.tipo.lower().strip(),
        descrizione=payload.descrizione or "",
        importo=float(payload.importo),
        data=payload.data or datetime.now().strftime("%Y-%m-%d"),
        fornitore_id=payload.fornitore_id,
        fornitore_nome=nome,
        documento_ref=payload.documento_ref or "",
        note=payload.note or "",
    )
    await db.spese_accessorie.insert_one(sp.model_dump())
    return sp


@router.put("/spese/{sid}", response_model=SpesaAccessoria)
async def update_spesa(sid: str, payload: SpesaCreate):
    existing = await db.spese_accessorie.find_one({"id": sid}, {"_id": 0})
    if not existing:
        raise HTTPException(404, "Spesa non trovata")
    data = {**existing, **payload.model_dump(exclude_unset=True)}
    sp = SpesaAccessoria(**_to_dt(data))
    await db.spese_accessorie.update_one({"id": sid}, {"$set": sp.model_dump()})
    return sp


@router.delete("/spese/{sid}")
async def delete_spesa(sid: str):
    r = await db.spese_accessorie.delete_one({"id": sid})
    if r.deleted_count == 0:
        raise HTTPException(404, "Spesa non trovata")
    return {"ok": True}


@router.get("/spese-report")
async def spese_report(anno: Optional[int] = None):
    """Aggregato spese per tipo (con totali). Filtro opzionale per anno."""
    match = {}
    if anno:
        match["data"] = {"$regex": f"^{anno}-"}
    pipeline = [
        {"$match": match} if match else {"$match": {}},
        {"$group": {
            "_id": "$tipo",
            "totale": {"$sum": "$importo"},
            "count": {"$sum": 1},
        }},
        {"$sort": {"totale": -1}},
    ]
    docs = await db.spese_accessorie.aggregate(pipeline).to_list(500)
    per_tipo = [{"tipo": d["_id"] or "altro", "totale": float(d["totale"]), "count": d["count"]} for d in docs]
    totale_generale = sum(x["totale"] for x in per_tipo)
    return {"anno": anno, "totale_generale": totale_generale, "per_tipo": per_tipo}


# ---------------------------------------------------------------------------
# EXPORT: Listino PDF e Inventario Excel
# ---------------------------------------------------------------------------

@router.get("/listino.pdf")
async def listino_pdf(
    fornitore_id: Optional[str] = None,
    categoria: Optional[str] = None,
):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    query = {}
    if fornitore_id:
        query["fornitore_id"] = fornitore_id
    if categoria:
        query["categoria"] = categoria
    docs = await db.articoli.find(query, {"_id": 0}).to_list(5000)
    # Ordina alfabeticamente per la descrizione visualizzata (fallback su nome)
    docs.sort(key=lambda d: ((d.get("descrizione") or d.get("nome") or "").strip().lower()))
    cantiere = await db.cantiere.find_one({"id": "default"}, {"_id": 0}) or {}
    fornitore = None
    if fornitore_id:
        fornitore = await db.fornitori.find_one({"id": fornitore_id}, {"_id": 0})

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    story = []

    nome_cantiere = cantiere.get("nome") or "Portomare"
    story.append(Paragraph(f"<b>{nome_cantiere}</b>", styles["Title"]))
    subtitle = "Listino accessori nautici"
    if fornitore:
        subtitle += f" — Fornitore: {fornitore.get('nome')}"
    if categoria:
        subtitle += f" — Categoria: {categoria}"
    story.append(Paragraph(subtitle, styles["Heading3"]))
    story.append(Paragraph(datetime.now().strftime("Aggiornato al %d/%m/%Y"), styles["Normal"]))
    story.append(Spacer(1, 8))

    # Mappa fornitori per sigla (fallback su nome se sigla non impostata)
    fornitori_list = await db.fornitori.find({}, {"_id": 0, "id": 1, "nome": 1, "abbreviazione": 1}).to_list(1000)
    forn_map = {f["id"]: (f.get("abbreviazione") or f.get("nome", "")) for f in fornitori_list}

    headers = ["Codice", "Forn.", "Descrizione", "U.M.", "Prezzo € (IVA inc.)"]
    data = [headers]
    for d in docs:
        # Una sola descrizione: preferisci il campo descrizione, altrimenti usa il nome
        testo = (d.get("descrizione") or d.get("nome") or "").strip()
        data.append([
            d.get("codice", "") or "—",
            forn_map.get(d.get("fornitore_id"), "") or "—",
            testo[:110],
            d.get("unita_misura", "pz") or "pz",
            f"{float(d.get('prezzo_listino', 0)) * 1.22:.2f}",
        ])
    table = Table(data, colWidths=[22*mm, 14*mm, 105*mm, 12*mm, 27*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (-1, 1), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)

    doc.build(story)
    buf.seek(0)
    fname = "listino_accessori.pdf"
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/ordine-fornitore.pdf")
async def ordine_fornitore_pdf(fornitore_id: Optional[str] = None):
    """Genera un PDF ordine con tutti gli articoli sotto scorta minima.
    Se fornitore_id è specificato limita all'ordine per quel fornitore."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    query = {}
    if fornitore_id:
        query["fornitore_id"] = fornitore_id
    tutti = await db.articoli.find(query, {"_id": 0}).sort([("categoria", 1), ("nome", 1)]).to_list(5000)
    # Solo articoli sotto scorta minima
    docs = [a for a in tutti if float(a.get("quantita", 0)) <= float(a.get("scorta_minima", 0))]

    cantiere = await db.cantiere.find_one({"id": "default"}, {"_id": 0}) or {}
    fornitore = None
    if fornitore_id:
        fornitore = await db.fornitori.find_one({"id": fornitore_id}, {"_id": 0})

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    story = []

    # Intestazione cantiere
    nome_cantiere = cantiere.get("nome") or "Portomare"
    story.append(Paragraph(f"<b>{nome_cantiere}</b>", styles["Title"]))
    riga_ind = " · ".join(x for x in [
        cantiere.get("indirizzo"),
        f"{cantiere.get('cap','')} {cantiere.get('citta','')} ({cantiere.get('provincia','')})".strip(),
        f"Tel {cantiere.get('telefono')}" if cantiere.get("telefono") else "",
        f"P.IVA {cantiere.get('piva')}" if cantiere.get("piva") else "",
    ] if x and x.strip() and x.strip() != "()")
    if riga_ind:
        story.append(Paragraph(riga_ind, styles["Normal"]))
    story.append(Spacer(1, 8))

    story.append(Paragraph("<b>ORDINE MATERIALE — Riassortimento scorte</b>", styles["Heading2"]))
    story.append(Paragraph(datetime.now().strftime("Data: %d/%m/%Y"), styles["Normal"]))

    if fornitore:
        story.append(Spacer(1, 6))
        story.append(Paragraph(f"<b>Spett.le {fornitore.get('nome')}</b>", styles["Heading3"]))
        info = []
        if fornitore.get("referente"): info.append(f"c.a. {fornitore['referente']}")
        if fornitore.get("indirizzo"): info.append(fornitore["indirizzo"])
        if fornitore.get("email"): info.append(fornitore["email"])
        if fornitore.get("telefono"): info.append(f"Tel {fornitore['telefono']}")
        if info:
            story.append(Paragraph(" · ".join(info), styles["Normal"]))
    story.append(Spacer(1, 10))

    if not docs:
        story.append(Paragraph(
            "<i>Nessun articolo attualmente sotto scorta minima.</i>", styles["Normal"]))
    else:
        story.append(Paragraph(
            f"Si richiede la fornitura del seguente materiale ({len(docs)} articoli sotto scorta):",
            styles["Normal"]))
        story.append(Spacer(1, 6))

        headers = ["Codice", "Descrizione", "U.M.", "Giacenza", "Scorta min.", "Q.tà da ordinare"]
        data = [headers]
        for d in docs:
            qt = float(d.get("quantita", 0))
            sm = float(d.get("scorta_minima", 0))
            # Suggerisce il doppio della scorta minima meno la giacenza attuale (arrotondato per eccesso)
            da_ordinare = max(int(sm * 2 - qt + 0.999), int(sm) or 1)
            data.append([
                d.get("codice", "") or "—",
                (d.get("nome", "") or "") + ((" — " + d.get("descrizione", "")) if d.get("descrizione") else ""),
                d.get("unita_misura", "pz") or "pz",
                f"{qt:g}",
                f"{sm:g}",
                str(da_ordinare),
            ])
        table = Table(data, colWidths=[25*mm, 75*mm, 15*mm, 20*mm, 20*mm, 30*mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            # Evidenziazione colonna "Q.tà da ordinare"
            ("BACKGROUND", (5, 1), (5, -1), colors.HexColor("#fef3c7")),
            ("FONTNAME", (5, 1), (5, -1), "Helvetica-Bold"),
        ]))
        story.append(table)

    story.append(Spacer(1, 14))
    story.append(Paragraph(
        "In attesa di conferma d'ordine con tempi di consegna, si porgono cordiali saluti.",
        styles["Normal"]))
    story.append(Spacer(1, 20))
    story.append(Paragraph(f"<b>{nome_cantiere}</b>", styles["Normal"]))
    story.append(Paragraph("_______________________________", styles["Normal"]))

    doc.build(story)
    buf.seek(0)
    fname = f"ordine_{(fornitore or {}).get('nome','fornitore').lower().replace(' ','_')}.pdf" if fornitore else "ordine_fornitore.pdf"
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/inventario.pdf")
async def inventario_pdf():
    """Export inventario in PDF A4 (Codice, Fornitore, Nome, U.M., Quantità, Prezzo acquisto, Valore giacenza)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    articoli = await db.articoli.find({}, {"_id": 0}).sort("nome", 1).to_list(5000)
    fornitori = await db.fornitori.find({}, {"_id": 0}).to_list(1000)
    forn_map = {f["id"]: (f.get("abbreviazione") or f.get("nome") or "") for f in fornitori}
    cantiere = await db.cantiere.find_one({"id": "default"}, {"_id": 0}) or {}

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=12*mm, rightMargin=12*mm, topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    story = []

    nome_cantiere = cantiere.get("nome") or "Portomare"
    story.append(Paragraph(f"<b>{nome_cantiere}</b>", styles["Title"]))
    story.append(Paragraph("Inventario magazzino", styles["Heading3"]))
    story.append(Paragraph(datetime.now().strftime("Aggiornato al %d/%m/%Y"), styles["Normal"]))
    story.append(Spacer(1, 8))

    headers = ["Codice", "Forn.", "Nome", "U.M.", "Q.tà", "P.acq. €", "Valore €"]
    data = [headers]
    tot_valore = 0.0
    for a in articoli:
        q = float(a.get("quantita", 0) or 0)
        p = float(a.get("prezzo_acquisto", 0) or 0)
        val = q * p
        tot_valore += val
        data.append([
            a.get("codice", "") or "—",
            forn_map.get(a.get("fornitore_id"), "") or "—",
            (a.get("nome", "") or "")[:80],
            a.get("unita_misura", "pz") or "pz",
            f"{q:g}",
            f"{p:.2f}",
            f"{val:.2f}",
        ])
    # Riga totale
    data.append(["TOTALE", "", "", "", "", "", f"{tot_valore:.2f}"])

    # Larghezze colonne minime, max spazio a "Nome". Somma = 186mm (A4 - 24mm margini)
    table = Table(data, colWidths=[20*mm, 12*mm, 95*mm, 10*mm, 12*mm, 17*mm, 20*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f8fafc")]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        # Riga totale evidenziata
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#fef3c7")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("SPAN", (0, -1), (5, -1)),
        ("ALIGN", (0, -1), (5, -1), "RIGHT"),
    ]))
    story.append(table)

    doc.build(story)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="inventario_magazzino.pdf"'},
    )


@router.get("/inventario.xlsx")
async def inventario_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    articoli = await db.articoli.find({}, {"_id": 0}).sort("nome", 1).to_list(5000)
    fornitori = await db.fornitori.find({}, {"_id": 0}).to_list(1000)
    forn_map = {f["id"]: (f.get("abbreviazione") or f.get("nome") or "") for f in fornitori}

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"

    headers = ["Codice", "Fornitore", "Nome", "U.M.", "Quantità",
               "Prezzo acquisto €", "Valore giacenza €"]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F172A")
    for col in range(1, len(headers)+1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    tot_valore = 0.0
    for a in articoli:
        q = float(a.get("quantita", 0))
        p = float(a.get("prezzo_acquisto", 0))
        val = q * p
        tot_valore += val
        ws.append([
            a.get("codice", ""),
            forn_map.get(a.get("fornitore_id"), ""),
            a.get("nome", ""),
            a.get("unita_misura", "pz"),
            q,
            p,
            val,
        ])

    last = ws.max_row + 1
    ws.cell(row=last, column=6, value="TOTALE VALORE").font = Font(bold=True)
    ws.cell(row=last, column=7, value=tot_valore).font = Font(bold=True)
    ws.cell(row=last, column=7).number_format = '#,##0.00 "€"'

    for i in (6, 7):
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=i, max_col=i):
            for cell in row:
                cell.number_format = '#,##0.00 "€"'

    widths = [14, 22, 34, 8, 10, 16, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="inventario_magazzino.xlsx"'},
    )
